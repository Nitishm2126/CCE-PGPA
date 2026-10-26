import io
import uuid
import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.security import create_access_token
from app.db.session import SessionLocal
from app.models.achievement import Achievement, AchievementStatus
from app.models.assignment import MentorAssignment
from app.models.component import PGPAComponent
from app.models.semester import Semester
from app.models.user import HodProfile, MentorProfile, StudentProfile, User, UserRole
from app.services.export_service import sanitize_for_spreadsheet


@pytest.fixture
def db_session():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def test_formula_injection_sanitization():
    """
    Verify spreadsheet formula injection (CWE-1236) protection.
    Values starting with =, +, -, @, \\t, \\r must be prefixed with a single quote.
    """
    assert sanitize_for_spreadsheet("=SUM(A1:B10)") == "'=SUM(A1:B10)"
    assert sanitize_for_spreadsheet("+cmd|' /C calc'!A0") == "'+cmd|' /C calc'!A0"
    assert sanitize_for_spreadsheet("-1+2") == "'-1+2"
    assert sanitize_for_spreadsheet("@calc") == "'@calc"
    assert sanitize_for_spreadsheet("\talert(1)") == "'\talert(1)"
    assert sanitize_for_spreadsheet("\rmalicious") == "'\rmalicious"

    # Benign text and numbers must remain unchanged
    assert sanitize_for_spreadsheet("Aditya Sharma") == "Aditya Sharma"
    assert sanitize_for_spreadsheet(9.85) == 9.85
    assert sanitize_for_spreadsheet(10) == 10
    assert sanitize_for_spreadsheet(True) is True
    assert sanitize_for_spreadsheet(None) == ""


def test_mentor_export_excel_success(client: TestClient, db_session: Session):
    """
    Verify authenticated mentor can export Excel workbook with all 4 worksheets.
    """
    mentor = db_session.query(User).filter(User.role == UserRole.MENTOR).first()
    assert mentor is not None, "Mentor must exist"

    mentor_token = create_access_token(subject=str(mentor.id), role=mentor.role.value)
    headers = {"Authorization": f"Bearer {mentor_token}"}

    res = client.get("/api/v1/mentor/export/excel", headers=headers)
    assert res.status_code == 200, f"Expected 200, got {res.status_code}: {res.text}"
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["content-type"]
    assert "attachment" in res.headers["content-disposition"]
    assert ".xlsx" in res.headers["content-disposition"]

    # Verify openpyxl can load the generated workbook
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    expected_sheets = ["Student Summary", "Semester Results", "Component Breakdown", "Submissions & Approvals"]
    for sheet_name in expected_sheets:
        assert sheet_name in wb.sheetnames, f"Missing sheet: {sheet_name}"

    summary_sheet = wb["Student Summary"]
    assert summary_sheet.cell(row=5, column=1).value == "S.No"
    assert summary_sheet.cell(row=5, column=2).value == "Register Number"
    assert summary_sheet.cell(row=5, column=3).value == "Student Name"


def test_mentor_export_pdf_success(client: TestClient, db_session: Session):
    """
    Verify authenticated mentor can export landscape PDF report with ReportLab.
    """
    mentor = db_session.query(User).filter(User.role == UserRole.MENTOR).first()
    assert mentor is not None, "Mentor must exist"

    mentor_token = create_access_token(subject=str(mentor.id), role=mentor.role.value)
    headers = {"Authorization": f"Bearer {mentor_token}"}

    res = client.get("/api/v1/mentor/export/pdf", headers=headers)
    assert res.status_code == 200, f"Expected 200, got {res.status_code}: {res.text}"
    assert "application/pdf" in res.headers["content-type"]
    assert "attachment" in res.headers["content-disposition"]
    assert ".pdf" in res.headers["content-disposition"]
    assert res.content.startswith(b"%PDF-"), "Response content must be a valid PDF file"


def test_hod_export_excel_and_pdf_success(client: TestClient, db_session: Session):
    """
    Verify authenticated HOD can export department-wide Excel and PDF reports.
    """
    hod = db_session.query(User).filter(User.role == UserRole.HOD).first()
    assert hod is not None, "HOD must exist"

    hod_token = create_access_token(subject=str(hod.id), role=hod.role.value)
    headers = {"Authorization": f"Bearer {hod_token}"}

    # 1. Excel export
    res_excel = client.get("/api/v1/hod/export/excel", headers=headers)
    assert res_excel.status_code == 200, f"Expected 200, got {res_excel.status_code}"
    assert ".xlsx" in res_excel.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(res_excel.content))
    assert len(wb.sheetnames) == 4

    # 2. PDF export
    res_pdf = client.get("/api/v1/hod/export/pdf", headers=headers)
    assert res_pdf.status_code == 200, f"Expected 200, got {res_pdf.status_code}"
    assert res_pdf.content.startswith(b"%PDF-")


def test_individual_student_export_filter(client: TestClient, db_session: Session):
    """
    Verify individual student filter exports report strictly containing that student.
    """
    mentor = db_session.query(User).filter(User.role == UserRole.MENTOR).first()
    assert mentor is not None

    assignment = (
        db_session.query(MentorAssignment)
        .filter(MentorAssignment.mentor_id == mentor.id, MentorAssignment.is_active == True)
        .first()
    )
    assert assignment is not None, "Mentor must have at least one active assignment"
    target_student_id = str(assignment.student_id)

    mentor_token = create_access_token(subject=str(mentor.id), role=mentor.role.value)
    headers = {"Authorization": f"Bearer {mentor_token}"}

    # Query with student_id filter
    res = client.get(f"/api/v1/mentor/export/excel?student_id={target_student_id}", headers=headers)
    assert res.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    summary_sheet = wb["Student Summary"]
    # Only 1 student data row (row 6 onwards)
    data_rows = [row for row in summary_sheet.iter_rows(min_row=6, values_only=True) if any(row)]
    assert len(data_rows) == 1, f"Expected exactly 1 student record, got {len(data_rows)}"


def test_mentor_cross_mentee_access_forbidden(client: TestClient, db_session: Session):
    """
    Verify a mentor CANNOT export data for a student assigned to a different mentor (403 Forbidden).
    """
    mentors = db_session.query(User).filter(User.role == UserRole.MENTOR).all()
    if len(mentors) < 2:
        pytest.skip("Requires at least two mentors to test cross-mentee access")

    mentor_a = mentors[0]
    mentor_b = mentors[1]

    # Find a student assigned strictly to mentor_b
    assign_b = (
        db_session.query(MentorAssignment)
        .filter(MentorAssignment.mentor_id == mentor_b.id, MentorAssignment.is_active == True)
        .first()
    )
    assert assign_b is not None

    mentor_a_token = create_access_token(subject=str(mentor_a.id), role=mentor_a.role.value)
    headers_a = {"Authorization": f"Bearer {mentor_a_token}"}

    # Mentor A tries to export Mentor B's student
    res = client.get(f"/api/v1/mentor/export/excel?student_id={assign_b.student_id}", headers=headers_a)
    assert res.status_code == 403, f"Expected 403 Forbidden, got {res.status_code}: {res.text}"
    assert "not assigned to your mentor account" in res.json()["detail"]


def test_hod_cross_department_access_forbidden(client: TestClient, db_session: Session):
    """
    Verify HOD cannot export data for students outside their department (403 Forbidden).
    """
    hod = db_session.query(User).filter(User.role == UserRole.HOD).first()
    assert hod is not None

    # Create temporary student from different department
    other_user = User(
        email=f"foreign_student_{uuid.uuid4().hex[:6]}@test.edu",
        full_name="Foreign Department Student",
        password_hash="test_argon2_password_hash",
        role=UserRole.STUDENT,
        is_active=True,
    )
    db_session.add(other_user)
    db_session.flush()

    other_profile = StudentProfile(
        user_id=other_user.id,
        register_number=f"REG{uuid.uuid4().hex[:6].upper()}",
        department="ECE",  # Not CCE
        batch="2024-2028",
        semester="IV",
    )
    db_session.add(other_profile)
    db_session.commit()

    try:
        hod_token = create_access_token(subject=str(hod.id), role=hod.role.value)
        headers = {"Authorization": f"Bearer {hod_token}"}

        res = client.get(f"/api/v1/hod/export/excel?student_id={other_user.id}", headers=headers)
        assert res.status_code == 403, f"Expected 403 Forbidden, got {res.status_code}: {res.text}"
        assert "does not belong to the" in res.json()["detail"]
    finally:
        db_session.delete(other_profile)
        db_session.delete(other_user)
        db_session.commit()


def test_export_unauthorized_and_student_forbidden(client: TestClient, db_session: Session):
    """
    Verify unauthenticated users get 401 and students get 403 on export endpoints.
    """
    # 1. No token -> 401
    assert client.get("/api/v1/mentor/export/excel").status_code == 401
    assert client.get("/api/v1/mentor/export/pdf").status_code == 401
    assert client.get("/api/v1/hod/export/excel").status_code == 401
    assert client.get("/api/v1/hod/export/pdf").status_code == 401

    # 2. Student token -> 403
    student = db_session.query(User).filter(User.role == UserRole.STUDENT).first()
    assert student is not None
    student_token = create_access_token(subject=str(student.id), role=student.role.value)
    headers = {"Authorization": f"Bearer {student_token}"}

    assert client.get("/api/v1/mentor/export/excel", headers=headers).status_code == 403
    assert client.get("/api/v1/mentor/export/pdf", headers=headers).status_code == 403
    assert client.get("/api/v1/hod/export/excel", headers=headers).status_code == 403
    assert client.get("/api/v1/hod/export/pdf", headers=headers).status_code == 403


def test_export_is_strictly_read_only(client: TestClient, db_session: Session):
    """
    Verify report export is strictly read-only and never mutates achievement statuses or logs.
    """
    mentor = db_session.query(User).filter(User.role == UserRole.MENTOR).first()
    assert mentor is not None

    mentor_token = create_access_token(subject=str(mentor.id), role=mentor.role.value)
    headers = {"Authorization": f"Bearer {mentor_token}"}

    # Count achievements and statuses prior to export
    achievements_before = db_session.query(Achievement.id, Achievement.status, Achievement.updated_at).all()

    # Perform both Excel and PDF exports
    client.get("/api/v1/mentor/export/excel", headers=headers)
    client.get("/api/v1/mentor/export/pdf", headers=headers)

    # Count achievements and statuses after export
    achievements_after = db_session.query(Achievement.id, Achievement.status, Achievement.updated_at).all()

    assert len(achievements_before) == len(achievements_after), "Achievement count must not change"
    assert achievements_before == achievements_after, "Achievement records and statuses must remain completely unchanged"
