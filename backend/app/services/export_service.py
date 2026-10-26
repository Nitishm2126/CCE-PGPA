"""
CCE-PGPA Institutional Report Export Service.
Generates comprehensive Excel (.xlsx) workbooks and PDF (.pdf) documents
for Faculty Mentors and Department HODs using openpyxl and ReportLab.

Security & Integrity Constraints:
- Enforces strict read-only execution: never mutates academic records or statuses.
- Enforces formula injection sanitization on all spreadsheet cells.
- Accurately renders derived letter grades, grade points, raw scores, credits, and PGPA tallies.
- Employs streaming BytesIO buffers to prevent server file-system debris.
"""

from datetime import datetime, timezone
import io
import re
from typing import Any, Dict, List, Optional
import uuid

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfgen import canvas
from reportlab.platypus import HRFlowable, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import func, select
from sqlalchemy.orm import Session, joinedload, selectinload

from app.models.achievement import Achievement, AchievementStatus
from app.models.assignment import MentorAssignment
from app.models.component import PGPAComponent
from app.models.grading import StudentComponentResult, StudentSemesterPGPA
from app.models.review import VerificationLog
from app.models.semester import Semester
from app.models.user import StudentProfile, User
from app.services.grading_engine import GRADE_POINTS_MAP, calculate_overall_pgpa


# ==============================================================================
# Security & Formatting Helpers
# ==============================================================================

def sanitize_for_spreadsheet(value: Any) -> Any:
    """
    Prevents CSV / Spreadsheet Formula Injection (CWE-1236).
    If a string starts with '=', '+', '-', '@', '\t', or '\r',
    prefixes with a single quote (') to force Excel to treat it as plain text.
    """
    if value is None:
        return ""
    if isinstance(value, str):
        if not value:
            return ""
        if value[0] in ("=", "+", "-", "@", "\t", "\r"):
            return f"'{value}"
        return value
    return value


class NumberedCanvas(canvas.Canvas):
    """
    Two-pass ReportLab canvas to compute and render total page count
    and formal institutional running headers/footers.
    """
    def __init__(self, *args: Any, **kwargs: Any) -> None:
        super().__init__(*args, **kwargs)
        self._saved_page_states: List[Dict[str, Any]] = []

    def showPage(self) -> None:
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self) -> None:
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count: int) -> None:
        self.saveState()
        page_width, page_height = landscape(A4)

        # Running Top Rule & Title
        self.setFont("Helvetica-Bold", 8)
        self.setFillColor(colors.HexColor("#1E3A8A"))
        self.drawString(40, page_height - 25, "SRI SAIRAM ENGINEERING COLLEGE — CCE ACHIEVE PGPA REPORT")

        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        self.drawRightString(page_width - 40, page_height - 25, "Confidential Academic Document")

        self.setStrokeColor(colors.HexColor("#E2E8F0"))
        self.setLineWidth(0.75)
        self.line(40, page_height - 28, page_width - 40, page_height - 28)

        # Running Bottom Rule & Footer
        self.line(40, 30, page_width - 40, 30)
        self.setFont("Helvetica", 7.5)
        self.drawString(40, 20, "Official Autonomous Academic Assessment • Verified by Faculty Mentor & Department HOD")
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(page_width - 40, 20, page_str)
        self.restoreState()


# ==============================================================================
# Report Data Builder
# ==============================================================================

def build_report_data(
    db: Session,
    student_ids: List[uuid.UUID],
    semester_id: Optional[uuid.UUID] = None,
    component_id: Optional[uuid.UUID] = None,
    status_filter: Optional[str] = None,
    department_name: str = "CCE",
    requester_role: str = "Faculty Mentor",
    requester_name: str = "Authorized Faculty",
) -> Dict[str, Any]:
    """
    Queries, aggregates, and organizes academic results, semester matrices,
    component breakdowns, and verification logs for authorized students.
    """
    now_utc = datetime.now(timezone.utc)

    # 1. Fetch Students & Profiles
    students: List[User] = []
    if student_ids:
        stmt = (
            select(User)
            .options(joinedload(User.student_profile))
            .where(User.id.in_(student_ids))
            .order_by(User.full_name.asc())
        )
        students = list(db.scalars(stmt).all())

    # 2. Fetch Active Mentor Assignments
    mentor_map: Dict[uuid.UUID, str] = {}
    if student_ids:
        assign_stmt = (
            select(MentorAssignment)
            .options(joinedload(MentorAssignment.mentor))
            .where(
                MentorAssignment.student_id.in_(student_ids),
                MentorAssignment.is_active.is_(True),
            )
        )
        assignments = db.scalars(assign_stmt).all()
        for a in assignments:
            if a.mentor:
                mentor_map[a.student_id] = a.mentor.full_name

    # 3. Fetch Semesters & Components Catalog
    all_semesters = list(db.scalars(select(Semester).order_by(Semester.semester_number.asc())).all())
    semester_by_id = {s.id: s for s in all_semesters}
    semester_by_num = {s.semester_number: s for s in all_semesters}

    all_components = list(db.scalars(select(PGPAComponent).order_by(PGPAComponent.display_order.asc())).all())
    component_by_id = {c.id: c for c in all_components}

    selected_semester = semester_by_id.get(semester_id) if semester_id else None
    selected_component = component_by_id.get(component_id) if component_id else None

    # 4. Fetch StudentSemesterPGPA Records
    pgpa_records_map: Dict[tuple[uuid.UUID, uuid.UUID], StudentSemesterPGPA] = {}
    if student_ids:
        pgpa_stmt = select(StudentSemesterPGPA).where(StudentSemesterPGPA.student_id.in_(student_ids))
        if semester_id:
            pgpa_stmt = pgpa_stmt.where(StudentSemesterPGPA.semester_id == semester_id)
        for r in db.scalars(pgpa_stmt).all():
            pgpa_records_map[(r.student_id, r.semester_id)] = r

    # 5. Fetch StudentComponentResults
    comp_results_stmt = (
        select(StudentComponentResult)
        .options(
            joinedload(StudentComponentResult.component),
            joinedload(StudentComponentResult.semester),
        )
        .where(StudentComponentResult.student_id.in_(student_ids))
    )
    if semester_id:
        comp_results_stmt = comp_results_stmt.where(StudentComponentResult.semester_id == semester_id)
    if component_id:
        comp_results_stmt = comp_results_stmt.where(StudentComponentResult.component_id == component_id)
    if status_filter == "approved":
        comp_results_stmt = comp_results_stmt.where(StudentComponentResult.is_verified.is_(True))

    comp_results = list(db.scalars(comp_results_stmt.order_by(StudentComponentResult.calculated_at.desc())).all())

    # Group component results by (student_id, semester_id) for on-the-fly PGPA fallback
    comp_results_by_student_sem: Dict[tuple[uuid.UUID, uuid.UUID], List[StudentComponentResult]] = {}
    for cr in comp_results:
        key = (cr.student_id, cr.semester_id)
        comp_results_by_student_sem.setdefault(key, []).append(cr)

    # 6. Fetch Submissions (Achievements)
    ach_stmt = (
        select(Achievement)
        .options(
            joinedload(Achievement.component),
            joinedload(Achievement.semester),
            joinedload(Achievement.student).joinedload(User.student_profile),
            selectinload(Achievement.verification_logs).joinedload(VerificationLog.reviewer),
        )
        .where(
            Achievement.student_id.in_(student_ids),
            Achievement.status != AchievementStatus.DRAFT,
        )
    )
    if semester_id:
        ach_stmt = ach_stmt.where(Achievement.semester_id == semester_id)
    if component_id:
        ach_stmt = ach_stmt.where(Achievement.component_id == component_id)
    if status_filter == "approved":
        ach_stmt = ach_stmt.where(Achievement.status == AchievementStatus.APPROVED)
    elif status_filter == "pending":
        ach_stmt = ach_stmt.where(Achievement.status == AchievementStatus.PENDING)
    elif status_filter == "rejected":
        ach_stmt = ach_stmt.where(Achievement.status == AchievementStatus.REJECTED)

    submissions = list(db.scalars(ach_stmt.order_by(Achievement.submitted_at.desc().nulls_last())).all())

    # 7. Build Student Summaries & Semester Matrix
    students_summary: List[Dict[str, Any]] = []
    semester_matrix: List[Dict[str, Any]] = []

    for s in students:
        p = s.student_profile
        mentor_name = mentor_map.get(s.id, (p.original_mentor_name if p else "Unassigned"))

        total_attempted = 0
        total_earned = 0
        weighted_pgpa_sum = 0.0
        weighted_credits = 0
        sem_scores: Dict[int, Optional[float]] = {}
        sem_gpas: Dict[int, Optional[float]] = {}
        sem_credits: Dict[int, int] = {}

        for sem in all_semesters:
            sem_num = sem.semester_number
            rec = pgpa_records_map.get((s.id, sem.id))

            # If no cached record, check if we can compute from component results
            if not rec and (s.id, sem.id) in comp_results_by_student_sem:
                crs = comp_results_by_student_sem[(s.id, sem.id)]
                comp_dicts = []
                gpa_val = None
                for c in crs:
                    if c.component and c.component.slug == "academics-gpa" and c.is_verified:
                        gpa_val = c.raw_score
                    comp_dicts.append({
                        "component_id": str(c.component_id),
                        "component_slug": c.component.slug if c.component else "",
                        "grade": c.grade,
                        "grade_points": c.grade_points,
                        "credit_value": c.credit_value,
                        "is_verified": c.is_verified,
                    })
                calc = calculate_overall_pgpa(comp_dicts, semester_gpa=gpa_val)
                pgpa_val = calc["pgpa"]
                sem_attempted = calc["total_credits_attempted"]
                sem_earned = calc["total_credits_earned"]
                gpa_disp = calc["semester_gpa"]
            elif rec:
                pgpa_val = rec.pgpa
                sem_attempted = rec.total_credits_attempted
                sem_earned = rec.total_credits_earned
                gpa_disp = rec.gpa
            else:
                pgpa_val = None
                sem_attempted = 0
                sem_earned = 0
                gpa_disp = None

            sem_scores[sem_num] = pgpa_val
            sem_gpas[sem_num] = gpa_disp
            sem_credits[sem_num] = sem_earned

            total_attempted += sem_attempted
            total_earned += sem_earned
            if pgpa_val is not None and sem_attempted > 0:
                weighted_pgpa_sum += pgpa_val * sem_attempted
                weighted_credits += sem_attempted

        cumulative_pgpa = (
            round(weighted_pgpa_sum / weighted_credits, 2)
            if weighted_credits > 0
            else (
                round(sum(v for v in sem_scores.values() if v is not None) / len([v for v in sem_scores.values() if v is not None]), 2)
                if any(v is not None for v in sem_scores.values())
                else None
            )
        )

        students_summary.append({
            "student_id": str(s.id),
            "register_number": p.register_number if p else "-",
            "full_name": s.full_name,
            "department": p.department if p else department_name,
            "year": p.year if p else 2,
            "section": p.section if p else "-",
            "mentor_name": mentor_name,
            "total_attempted": total_attempted,
            "total_earned": total_earned,
            "cumulative_pgpa": cumulative_pgpa,
        })

        semester_matrix.append({
            "student_id": str(s.id),
            "register_number": p.register_number if p else "-",
            "full_name": s.full_name,
            "sem_scores": sem_scores,
            "sem_gpas": sem_gpas,
            "cumulative_pgpa": cumulative_pgpa,
        })

    # 8. Build Component Breakdown Items
    student_lookup = {s.id: s for s in students}
    component_breakdown_items: List[Dict[str, Any]] = []
    for cr in comp_results:
        st = student_lookup.get(cr.student_id)
        if not st:
            continue
        p = st.student_profile
        component_breakdown_items.append({
            "register_number": p.register_number if p else "-",
            "student_name": st.full_name,
            "component_name": cr.component.name if cr.component else "-",
            "component_slug": cr.component.slug if cr.component else "-",
            "semester_name": cr.semester.display_name if cr.semester else f"Sem {cr.semester_id}",
            "semester_number": cr.semester.semester_number if cr.semester else 0,
            "raw_score": cr.raw_score if cr.raw_score is not None else "-",
            "grade": cr.grade,
            "grade_points": cr.grade_points,
            "credit_value": cr.credit_value,
            "earned_points": cr.grade_points * cr.credit_value,
            "is_verified": cr.is_verified,
            "source": cr.source,
        })

    # 9. Build Submissions Items
    submissions_items: List[Dict[str, Any]] = []
    for ach in submissions:
        p = ach.student.student_profile if ach.student else None
        latest_rev = ach.verification_logs[-1] if ach.verification_logs else None
        submissions_items.append({
            "register_number": p.register_number if p else "-",
            "student_name": ach.student.full_name if ach.student else "-",
            "title": ach.title,
            "component_name": ach.component.name if ach.component else "-",
            "semester_name": ach.semester.display_name if ach.semester else "-",
            "status": ach.status.value.upper(),
            "submitted_at": ach.submitted_at.strftime("%Y-%m-%d %H:%M") if ach.submitted_at else (
                ach.created_at.strftime("%Y-%m-%d %H:%M") if ach.created_at else "-"
            ),
            "reviewer_name": latest_rev.reviewer.full_name if (latest_rev and latest_rev.reviewer) else "-",
            "reviewed_at": latest_rev.created_at.strftime("%Y-%m-%d %H:%M") if latest_rev else "-",
            "remarks": latest_rev.remarks if latest_rev else "-",
        })

    # Scope and filter descriptions
    is_single_student = len(students) == 1
    scope_title = (
        f"Individual Student Report: {students[0].full_name} ({students[0].student_profile.register_number})"
        if is_single_student and students[0].student_profile
        else (f"All Assigned Mentees ({len(students)})" if requester_role.lower() == "mentor" else f"Department Roster ({len(students)})")
    )

    filter_desc = []
    if selected_semester:
        filter_desc.append(f"Semester: {selected_semester.display_name}")
    else:
        filter_desc.append("Semesters: All (1 to 8)")

    if selected_component:
        filter_desc.append(f"Component: {selected_component.name}")
    else:
        filter_desc.append("Components: All Official (14)")

    if status_filter:
        filter_desc.append(f"Status: {status_filter.capitalize()}")

    return {
        "metadata": {
            "department": department_name,
            "generated_at": now_utc.strftime("%Y-%m-%d %H:%M:%S UTC"),
            "date_stamp": now_utc.strftime("%Y%m%d_%H%M%S"),
            "requester_role": requester_role,
            "requester_name": requester_name,
            "scope_title": scope_title,
            "filter_summary": " • ".join(filter_desc),
            "student_count": len(students),
            "is_single_student": is_single_student,
        },
        "all_semesters": [{"number": sem.semester_number, "name": sem.display_name} for sem in all_semesters],
        "students_summary": students_summary,
        "semester_matrix": semester_matrix,
        "component_breakdown": component_breakdown_items,
        "submissions_list": submissions_items,
    }


# ==============================================================================
# Excel Report Generator (openpyxl)
# ==============================================================================

def generate_excel_report(report_data: Dict[str, Any]) -> bytes:
    """
    Constructs a professional, multi-worksheet Excel workbook.
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    meta = report_data["metadata"]
    dept = meta.get("department", "CCE")

    # Typography & Styles
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    navy_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    dark_slate_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    accent_sub_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    def apply_title_header(ws: Any, sheet_title: str, max_col: int) -> None:
        last_col_letter = get_column_letter(max_col)
        # Institution Banner
        ws.merge_cells(f"A1:{last_col_letter}1")
        c1 = ws["A1"]
        c1.value = f"SRI SAIRAM ENGINEERING COLLEGE — {dept} DEPARTMENT"
        c1.font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
        c1.fill = dark_slate_fill
        c1.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        # Sheet Title Banner
        ws.merge_cells(f"A2:{last_col_letter}2")
        c2 = ws["A2"]
        c2.value = f"CCE ACHIEVE PGPA: {sheet_title.upper()} — {meta['scope_title']}"
        c2.font = Font(name="Segoe UI", size=11, bold=True, color="1E3A8A")
        c2.fill = accent_sub_fill
        c2.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 22

        # Filter & Metadata line
        ws.merge_cells(f"A3:{last_col_letter}3")
        c3 = ws["A3"]
        c3.value = f"Filters: {meta['filter_summary']} | Generated by: {meta['requester_name']} ({meta['requester_role']}) on {meta['generated_at']}"
        c3.font = Font(name="Segoe UI", size=9, italic=True, color="475569")
        c3.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        c3.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[3].height = 20

    # --------------------------------------------------------------------------
    # Tab 1: Student Summary
    # --------------------------------------------------------------------------
    ws1 = wb.create_sheet(title="Student Summary")
    summary_headers = [
        "S.No", "Register Number", "Student Name", "Department", "Year / Sec",
        "Assigned Mentor", "Attempted Credits", "Earned Credits", "Cumulative PGPA"
    ]
    apply_title_header(ws1, "Academic Performance Summary", len(summary_headers))

    # Header Row
    ws1.row_dimensions[5].height = 25
    for col_idx, h in enumerate(summary_headers, 1):
        cell = ws1.cell(row=5, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    row_num = 6
    for idx, s in enumerate(report_data["students_summary"], 1):
        ws1.row_dimensions[row_num].height = 20
        fill = zebra_fill if idx % 2 == 0 else white_fill
        vals = [
            idx,
            sanitize_for_spreadsheet(s["register_number"]),
            sanitize_for_spreadsheet(s["full_name"]),
            sanitize_for_spreadsheet(s["department"]),
            f"Year {s['year']} - {s['section']}",
            sanitize_for_spreadsheet(s["mentor_name"]),
            s["total_attempted"],
            s["total_earned"],
            s["cumulative_pgpa"] if s["cumulative_pgpa"] is not None else "N/A",
        ]
        for col_idx, v in enumerate(vals, 1):
            cell = ws1.cell(row=row_num, column=col_idx, value=v)
            cell.font = Font(name="Segoe UI", size=9.5)
            cell.fill = fill
            cell.border = thin_border
            if col_idx in (1, 2, 5):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in (7, 8, 9):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        row_num += 1

    # --------------------------------------------------------------------------
    # Tab 2: Semester-wise PGPA Matrix
    # --------------------------------------------------------------------------
    ws2 = wb.create_sheet(title="Semester Results")
    sem_cols = [f"Sem {i}" for i in range(1, 9)]
    matrix_headers = ["S.No", "Register Number", "Student Name"] + sem_cols + ["Cumulative PGPA"]
    apply_title_header(ws2, "Eight Semester PGPA Matrix", len(matrix_headers))

    ws2.row_dimensions[5].height = 25
    for col_idx, h in enumerate(matrix_headers, 1):
        cell = ws2.cell(row=5, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    row_num = 6
    for idx, m in enumerate(report_data["semester_matrix"], 1):
        ws2.row_dimensions[row_num].height = 20
        fill = zebra_fill if idx % 2 == 0 else white_fill
        vals = [
            idx,
            sanitize_for_spreadsheet(m["register_number"]),
            sanitize_for_spreadsheet(m["full_name"]),
        ]
        for i in range(1, 9):
            sc = m["sem_scores"].get(i)
            vals.append(sc if sc is not None else "-")
        vals.append(m["cumulative_pgpa"] if m["cumulative_pgpa"] is not None else "N/A")

        for col_idx, v in enumerate(vals, 1):
            cell = ws2.cell(row=row_num, column=col_idx, value=v)
            cell.font = Font(name="Segoe UI", size=9.5)
            cell.fill = fill
            cell.border = thin_border
            if col_idx in (1, 2):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx >= 4:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        row_num += 1

    # --------------------------------------------------------------------------
    # Tab 3: Component Breakdown
    # --------------------------------------------------------------------------
    ws3 = wb.create_sheet(title="Component Breakdown")
    comp_headers = [
        "S.No", "Register Number", "Student Name", "Semester", "Component Name",
        "Raw Score", "Letter Grade", "Grade Points", "Credits", "Earned Points", "Status"
    ]
    apply_title_header(ws3, "14 Official PGPA Components Evaluation", len(comp_headers))

    ws3.row_dimensions[5].height = 25
    for col_idx, h in enumerate(comp_headers, 1):
        cell = ws3.cell(row=5, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    row_num = 6
    for idx, cb in enumerate(report_data["component_breakdown"], 1):
        ws3.row_dimensions[row_num].height = 20
        fill = zebra_fill if idx % 2 == 0 else white_fill
        vals = [
            idx,
            sanitize_for_spreadsheet(cb["register_number"]),
            sanitize_for_spreadsheet(cb["student_name"]),
            sanitize_for_spreadsheet(cb["semester_name"]),
            sanitize_for_spreadsheet(cb["component_name"]),
            sanitize_for_spreadsheet(str(cb["raw_score"])),
            sanitize_for_spreadsheet(cb["grade"]),
            cb["grade_points"],
            cb["credit_value"],
            cb["earned_points"],
            "Verified" if cb["is_verified"] else "Pending",
        ]
        for col_idx, v in enumerate(vals, 1):
            cell = ws3.cell(row=row_num, column=col_idx, value=v)
            cell.font = Font(name="Segoe UI", size=9.5)
            cell.fill = fill
            cell.border = thin_border
            if col_idx in (1, 2, 4, 7, 11):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in (8, 9, 10):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        row_num += 1

    # --------------------------------------------------------------------------
    # Tab 4: Submissions & Approvals
    # --------------------------------------------------------------------------
    ws4 = wb.create_sheet(title="Submissions & Approvals")
    sub_headers = [
        "S.No", "Register Number", "Student Name", "Activity Title", "Component",
        "Semester", "Approval Status", "Submitted Date", "Reviewer", "Review Date", "Remarks"
    ]
    apply_title_header(ws4, "Activity Submissions & Review Decisions", len(sub_headers))

    ws4.row_dimensions[5].height = 25
    for col_idx, h in enumerate(sub_headers, 1):
        cell = ws4.cell(row=5, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    row_num = 6
    for idx, sb in enumerate(report_data["submissions_list"], 1):
        ws4.row_dimensions[row_num].height = 20
        fill = zebra_fill if idx % 2 == 0 else white_fill
        vals = [
            idx,
            sanitize_for_spreadsheet(sb["register_number"]),
            sanitize_for_spreadsheet(sb["student_name"]),
            sanitize_for_spreadsheet(sb["title"]),
            sanitize_for_spreadsheet(sb["component_name"]),
            sanitize_for_spreadsheet(sb["semester_name"]),
            sanitize_for_spreadsheet(sb["status"]),
            sanitize_for_spreadsheet(sb["submitted_at"]),
            sanitize_for_spreadsheet(sb["reviewer_name"]),
            sanitize_for_spreadsheet(sb["reviewed_at"]),
            sanitize_for_spreadsheet(sb["remarks"]),
        ]
        for col_idx, v in enumerate(vals, 1):
            cell = ws4.cell(row=row_num, column=col_idx, value=v)
            cell.font = Font(name="Segoe UI", size=9.5)
            cell.fill = fill
            cell.border = thin_border
            if col_idx in (1, 2, 6, 7, 8, 10):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        row_num += 1

    # Auto-adjust column widths across all sheets
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # ignore merged title rows for column width calculation
                if cell.row in (1, 2, 3):
                    continue
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


# ==============================================================================
# PDF Report Generator (ReportLab)
# ==============================================================================

def generate_pdf_report(report_data: Dict[str, Any]) -> bytes:
    """
    Constructs an institutional, publication-quality PDF report using ReportLab.
    Uses landscape A4 with clean tables, wrap-safe cells, and running headers/footers.
    """
    meta = report_data["metadata"]
    dept = meta.get("department", "CCE")
    page_w, page_h = landscape(A4)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=36,
        rightMargin=36,
        topMargin=38,
        bottomMargin=38,
    )

    styles = getSampleStyleSheet()

    # Custom typography styles
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=14,
        leading=17,
        textColor=colors.HexColor("#1E3A8A"),
        alignment=0,
    )
    subtitle_style = ParagraphStyle(
        "ReportSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#0F172A"),
        alignment=0,
    )
    meta_style = ParagraphStyle(
        "ReportMeta",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#475569"),
    )
    section_title_style = ParagraphStyle(
        "SectionTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10.5,
        leading=14,
        textColor=colors.HexColor("#1E3A8A"),
        spaceAfter=4,
    )
    th_style = ParagraphStyle(
        "TableHead",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=7.5,
        leading=9,
        textColor=colors.white,
        alignment=1,
    )
    td_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#1E293B"),
    )
    td_center = ParagraphStyle(
        "TableCellCenter",
        parent=td_style,
        alignment=1,
    )
    td_bold = ParagraphStyle(
        "TableCellBold",
        parent=td_style,
        fontName="Helvetica-Bold",
        textColor=colors.HexColor("#0F172A"),
    )

    story: List[Any] = []

    # 1. Institutional Banner
    story.append(Paragraph(f"<b>SRI SAIRAM ENGINEERING COLLEGE</b> (Autonomous) • DEPARTMENT OF {dept}", title_style))
    story.append(Spacer(1, 2))
    story.append(Paragraph(f"CCE ACHIEVE PGPA EVALUATION REPORT: {meta['scope_title']}", subtitle_style))
    story.append(Spacer(1, 4))

    meta_text = (
        f"<b>Active Filters:</b> {meta['filter_summary']}<br/>"
        f"<b>Generated By:</b> {meta['requester_name']} ({meta['requester_role']}) &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<b>Timestamp:</b> {meta['generated_at']} &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<b>Total Cohort:</b> {meta['student_count']} Students"
    )
    story.append(Paragraph(meta_text, meta_style))
    story.append(Spacer(1, 8))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#CBD5E1"), spaceAfter=10))

    # 2. Section: Student Performance Summary
    story.append(Paragraph("1. Student Performance & Cumulative PGPA Summary", section_title_style))

    summary_headers = [
        Paragraph("S.No", th_style),
        Paragraph("Register No", th_style),
        Paragraph("Student Name", th_style),
        Paragraph("Dept / Year", th_style),
        Paragraph("Assigned Mentor", th_style),
        Paragraph("Attempted Cr", th_style),
        Paragraph("Earned Cr", th_style),
        Paragraph("Cumulative PGPA", th_style),
    ]

    summary_rows: List[List[Any]] = [summary_headers]
    for idx, s in enumerate(report_data["students_summary"], 1):
        summary_rows.append([
            Paragraph(str(idx), td_center),
            Paragraph(str(s["register_number"]), td_center),
            Paragraph(str(s["full_name"]), td_bold),
            Paragraph(f"{s['department']} - Y{s['year']}", td_center),
            Paragraph(str(s["mentor_name"]), td_style),
            Paragraph(str(s["total_attempted"]), td_center),
            Paragraph(str(s["total_earned"]), td_center),
            Paragraph(
                f"<b>{s['cumulative_pgpa']:.2f}</b>" if s["cumulative_pgpa"] is not None else "-",
                td_center,
            ),
        ])

    avail_w = page_w - 72  # 770 pt on landscape A4
    col_w_summary = [30, 80, 180, 70, 190, 70, 70, 80]  # sum = 770
    table_summary = Table(summary_rows, colWidths=col_w_summary, repeatRows=1)
    table_summary.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(table_summary)
    story.append(Spacer(1, 14))

    # 3. Section: Eight Semester Results Matrix
    story.append(Paragraph("2. Eight-Semester PGPA Matrix (Semesters 1 to 8)", section_title_style))

    matrix_headers = (
        [Paragraph("S.No", th_style), Paragraph("Register No", th_style), Paragraph("Student Name", th_style)]
        + [Paragraph(f"Sem {i}", th_style) for i in range(1, 9)]
        + [Paragraph("Overall", th_style)]
    )
    matrix_rows: List[List[Any]] = [matrix_headers]
    for idx, m in enumerate(report_data["semester_matrix"], 1):
        r = [
            Paragraph(str(idx), td_center),
            Paragraph(str(m["register_number"]), td_center),
            Paragraph(str(m["full_name"]), td_bold),
        ]
        for i in range(1, 9):
            sc = m["sem_scores"].get(i)
            r.append(Paragraph(f"{sc:.2f}" if sc is not None else "-", td_center))
        r.append(Paragraph(f"<b>{m['cumulative_pgpa']:.2f}</b>" if m["cumulative_pgpa"] is not None else "-", td_center))
        matrix_rows.append(r)

    col_w_matrix = [25, 75, 170] + [52] * 8 + [64]  # sum = 270 + 416 + 64 = 750 (within 770)
    table_matrix = Table(matrix_rows, colWidths=col_w_matrix, repeatRows=1)
    table_matrix.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(table_matrix)
    story.append(Spacer(1, 14))

    # 4. Section: Component Evaluation Breakdown (if any)
    if report_data["component_breakdown"]:
        story.append(PageBreak())
        story.append(Paragraph("3. Evaluated Component Details (14 Official PGPA Components)", section_title_style))

        comp_headers = [
            Paragraph("S.No", th_style),
            Paragraph("Register No", th_style),
            Paragraph("Student Name", th_style),
            Paragraph("Semester", th_style),
            Paragraph("Component Name", th_style),
            Paragraph("Score", th_style),
            Paragraph("Grade", th_style),
            Paragraph("Points", th_style),
            Paragraph("Credit", th_style),
            Paragraph("Status", th_style),
        ]
        comp_rows: List[List[Any]] = [comp_headers]
        for idx, cb in enumerate(report_data["component_breakdown"], 1):
            comp_rows.append([
                Paragraph(str(idx), td_center),
                Paragraph(str(cb["register_number"]), td_center),
                Paragraph(str(cb["student_name"]), td_style),
                Paragraph(str(cb["semester_name"]), td_center),
                Paragraph(str(cb["component_name"]), td_style),
                Paragraph(str(cb["raw_score"]), td_center),
                Paragraph(f"<b>{cb['grade']}</b>", td_center),
                Paragraph(str(cb["grade_points"]), td_center),
                Paragraph(str(cb["credit_value"]), td_center),
                Paragraph("Verified" if cb["is_verified"] else "Pending", td_center),
            ])

        col_w_comp = [25, 75, 140, 70, 190, 50, 45, 45, 45, 65]  # sum = 750
        table_comp = Table(comp_rows, colWidths=col_w_comp, repeatRows=1)
        table_comp.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(table_comp)
        story.append(Spacer(1, 14))

    # 5. Section: Submissions & Approvals Trail (if any)
    if report_data["submissions_list"]:
        story.append(Spacer(1, 10))
        story.append(Paragraph("4. Student Activity Submissions & Approval Decisions", section_title_style))

        sub_headers = [
            Paragraph("S.No", th_style),
            Paragraph("Register No", th_style),
            Paragraph("Activity Title", th_style),
            Paragraph("Component", th_style),
            Paragraph("Semester", th_style),
            Paragraph("Status", th_style),
            Paragraph("Submitted Date", th_style),
            Paragraph("Reviewer", th_style),
            Paragraph("Remarks", th_style),
        ]
        sub_rows: List[List[Any]] = [sub_headers]
        for idx, sb in enumerate(report_data["submissions_list"], 1):
            sub_rows.append([
                Paragraph(str(idx), td_center),
                Paragraph(str(sb["register_number"]), td_center),
                Paragraph(str(sb["title"]), td_style),
                Paragraph(str(sb["component_name"]), td_style),
                Paragraph(str(sb["semester_name"]), td_center),
                Paragraph(f"<b>{sb['status']}</b>", td_center),
                Paragraph(str(sb["submitted_at"]), td_center),
                Paragraph(str(sb["reviewer_name"]), td_style),
                Paragraph(str(sb["remarks"]), td_style),
            ])

        col_w_sub = [25, 75, 150, 130, 60, 65, 80, 85, 80]  # sum = 750
        table_sub = Table(sub_rows, colWidths=col_w_sub, repeatRows=1)
        table_sub.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(table_sub)

    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer.getvalue()
